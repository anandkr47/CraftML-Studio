"""Text extraction for knowledge-base documents."""

from io import BytesIO
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".md", ".markdown", ".txt", ".docx"}


class UnsupportedDocument(ValueError):
    pass


def extract_text(filename: str, data: bytes) -> str:
    """Return the plain text of a PDF, DOCX, Markdown or TXT file."""
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedDocument(f"Unsupported file type '{ext}'. Use PDF, DOCX, Markdown or TXT.")

    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(BytesIO(data))
        return "\n\n".join(page.extract_text() or "" for page in reader.pages)

    if ext == ".docx":
        from docx import Document

        return "\n\n".join(p.text for p in Document(BytesIO(data)).paragraphs if p.text.strip())

    return data.decode("utf-8", errors="replace")
